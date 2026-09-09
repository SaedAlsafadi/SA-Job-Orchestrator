"""Phase 19.5 CV-pipeline integrity tests (mocked AI - no credits spent).

Covers the full corrective acceptance path:

    corrupted base CV text
    -> structured document builder (normalize + parse)
    -> mock tailoring changes (accept/reject)
    -> deterministic merge
    -> PDF + DOCX rendering
    -> document integrity verification
"""

from types import SimpleNamespace

import pytest

from app.core.documents.generator import DocumentGenerator
from app.core.llm.prompts.resume_tailor import TailoredResumeData
from app.models.enums import ChangeType
from app.services.pdf_verifier import check_document_integrity, verify_pdf_document
from app.services.resume_text import (
    build_resume_data_from_text,
    normalize_resume_text,
)
from app.services.tailoring_merge import merge_tailoring_changes

# A representative corrupted base CV mirroring the observed real-world PDF
# extraction: letter-spaced headings, glued blocks, mojibake bullets, fused
# category names, and no extractable SKILLS header.
CORRUPTED_BASE_CV = "\n".join([
    "S A E D  T E S T",
    "A I  &  S o f t w a r e  E n g i n e e r",
    "Address: Riyadh, Saudi Arabia | Email: test@example.com | Phone: +966 50 123 4567",
    "S U M M A R Y",
    "AI Engineer with startup experience delivering AI-enabled business solutions and",
    "full-stack web applications. Skilled in modern web development.",
    "Managed international admission workflows across multiple institutions, digital",
    "platforms, and stakeholders. Te chnology Operations Specialist | ACME IT | Sep 2024 - Present WORK",
    "EXPERIENCE",
    "Research Data Specialist | Research Center | 1 Month Project",
    "Built and maintained a structured research publication database containing 400+",
    "standardized records for institutional reporting.",
    "Python",
    "SQL SKILL SAI & Machine Learning",
    "TensorFlow",
    "Keras",
    "EDUCATION",
    "Bachelor of Science in Computer Science | Imam University | FEB 2022 - JUNE 2026",
    "CERTIFICATIONS",
    "\u00f2 Elements of AI for Business | University of Helsinki",
    "P R O J E C T S :",
    "AI-Powered Platform (Daqeeq)",
    "Designed a full-stack B2B SaaS platform that leverages machine learning algorithms to",
    "forecast inventory demand.React \u2022 Next.js \u2022 Supabase \u2022 Machine Learning",
])

EXPECTED_SECTIONS = [
    "Experience", "Projects", "Skills", "Education", "Certifications",
]


def _mock_change(change_id: str, target: str, change_type: ChangeType, proposed: str) -> SimpleNamespace:
    return SimpleNamespace(
        change_id=change_id,
        target_reference=target,
        change_type=change_type,
        proposed_text=proposed,
        original_text="",
    )


def _docx_text(path) -> str:
    from docx import Document

    doc = Document(str(path))
    parts = [p.text for p in doc.paragraphs]
    for table in doc.tables:  # tables are separate containers in DOCX
        for row in table.rows:
            parts.extend(c.text for c in row.cells)
    return "\n".join(parts)


@pytest.mark.asyncio
async def test_base_section_preservation_from_corrupted_text():
    """All base sections survive normalization + structuring (PART B)."""
    from app.services.resume_text import build_resume_data_from_text

    d = build_resume_data_from_text(CORRUPTED_BASE_CV)
    assert d["name"] == "SAED TEST"
    assert d["title"] == "AI & Software Engineer"
    assert d["location"] == "Riyadh, Saudi Arabia"
    assert d["email"] == "test@example.com"
    assert d["experience"], "experience must be recovered"
    assert d["projects"], "projects must be recovered (legacy builder dropped them)"
    assert d["skills"], "skills must be recovered"
    assert d["education"], "education must be recovered"
    assert d["certifications"], "certifications must be recovered"
    assert d["summary"], "summary must be recovered"
    # Mojibake bullet normalized into the certification text.
    assert any("Elements of AI" in c for c in d["certifications"])
    assert all(not c.startswith("•") for c in d["certifications"])
    assert "demand. React" in d["projects"][0]["description"]


@pytest.mark.asyncio
async def test_mocked_tailoring_full_pipeline(tmp_path):
    """base -> mock changes -> accept/reject -> merge -> render -> verify (PART E)."""
    from app.core.documents.generator import DocumentGenerator
    from app.services.resume_text import build_resume_data_from_text

    base_dict = build_resume_data_from_text(CORRUPTED_BASE_CV)
    base_doc = TailoredResumeData.model_validate(base_dict)

    # The workbench decides acceptance upstream; only ACCEPTED changes reach
    # the merge. The rejected proposal must never appear in the artifact.
    accepted = _mock_change(
        "c1", "summary", ChangeType.MODIFY,
        "Tailored summary emphasizing AI procurement platforms and predictive analytics.",
    )
    rejected_text = "Invented skill that was rejected"
    rejected = _mock_change("c2", "skills", ChangeType.MODIFY, rejected_text)
    assert rejected.proposed_text == rejected_text  # simulated rejection path
    merged = merge_tailoring_changes(base_doc, [accepted])

    generator = DocumentGenerator(llm_client=None, output_dir=tmp_path)
    doc_res = await generator.generate_resume(
        resume_data=merged.model_dump(),
        job_description="",
        template_name="modern",
        formats=["pdf", "docx"],
    )
    assert doc_res.pdf_path, "PDF render must succeed"
    assert doc_res.docx_path, "DOCX render must succeed"

    text = _docx_text(doc_res.docx_path)
    integrity = check_document_integrity(
        text,
        expected_sections=[
            "Experience", "Projects", "Skills", "Education", "Certifications",
        ],
        required_texts=[
            "Tailored summary emphasizing AI procurement platforms",
            "AI-Powered Platform (Daqeeq)",
        ],
        forbidden_texts=[rejected_text],
        base_text=normalize_resume_text(CORRUPTED_BASE_CV),
    )
    assert integrity.is_valid, integrity.reason
    assert not integrity.missing_sections, integrity.missing_sections
    assert not integrity.forbidden_texts_found
    assert "demand.React" not in text

    pdf_integrity = verify_pdf_document(
        doc_res.pdf_path,
        expected_name=merged.name,
        expected_sections=EXPECTED_SECTIONS,
        required_texts=[
            "Tailored summary emphasizing AI procurement platforms",
            "AI-Powered Platform (Daqeeq)",
        ],
        forbidden_texts=[rejected_text],
        base_text=normalize_resume_text(CORRUPTED_BASE_CV),
    )
    assert pdf_integrity.is_valid, pdf_integrity.reason
    assert pdf_integrity.page_count > 0


# ---------------------------------------------------------------------------
# Verifier unit checks (PART D) - pure text, no PDF rendering needed
# ---------------------------------------------------------------------------


def test_integrity_flags_missing_section():
    result = check_document_integrity(
        "SAED TEST\nExperience\nBuilt things.",
        expected_sections=["Experience", "Projects", "Education"],
    )
    assert not result.is_valid
    assert "Projects" in result.missing_sections and "Education" in result.missing_sections


def test_integrity_flags_rejected_change_present():
    result = check_document_integrity(
        "summary text\nInvented skill that was rejected",
        forbidden_texts=["Invented skill that was rejected"],
    )
    assert not result.is_valid
    assert result.forbidden_texts_found


def test_integrity_flags_accepted_change_missing():
    result = check_document_integrity(
        "summary text only",
        required_texts=["Tailored summary emphasizing AI platforms"],
    )
    assert not result.is_valid
    assert result.missing_texts


def test_integrity_detects_concatenation_artifacts():
    # Glued ALL-CAPS header + glued sentence boundaries (observed defects).
    text = "\n".join([
        "experience description ending here." * 4,
        "InstitutesPROJECTS :",
        "predictive analytics.React stuff",
        "more lifecycle.Technology things here",
        "InstitutesSKILLS section content",
        "summary text.Profile content follows",
        "additional lines to pass the sparse check",
    ])
    result = check_document_integrity(text)
    assert not result.is_valid
    assert result.concatenation_issues


def test_integrity_flags_low_content_retention():
    base = ("Experienced engineer with deep expertise. " * 40).strip()
    result = check_document_integrity(
        "tiny doc",
        base_text=base,
        min_retention=0.5,
    )
    assert not result.is_valid
    assert "retention too low" in (result.reason or "")


def test_integrity_passes_clean_document():
    text = "\n".join([
        "SAED TEST", "AI & Software Engineer",
        "Professional Summary", "Tailored summary text.",
        "Experience", "Engineer — Acme (2020 - Now)", "Built things.",
        "Projects", "Daqeeq", "Procurement platform.",
        "Skills", "Python, FastAPI",
        "Education", "BSc — University (2022)",
        "Certifications", "AI for Business",
    ])
    result = check_document_integrity(
        text,
        expected_sections=EXPECTED_SECTIONS,
        required_texts=["Tailored summary text."],
        forbidden_texts=["fabricated"],
        base_text=text,
    )
    assert result.is_valid, result.reason
